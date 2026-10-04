"""Deterministic, wholly synthetic document builders for extraction tests.

These build real DOCX/PPTX/XLSX/PDF bytes with the standard libraries, never
fake extension bytes, and contain no real-student content.
"""
from __future__ import annotations

import io
from pathlib import Path
import zipfile


def build_docx(*, paragraphs=("First paragraph.", "Second paragraph."),
               heading="Synthetic Heading", table=None, tracked=None,
               tabs=False, bold_word=None) -> bytes:
    """Build a minimal but real DOCX with optional table, tabs, and revisions."""
    from docx import Document

    document = Document()
    if heading:
        document.add_heading(heading, level=1)
    for text in paragraphs:
        paragraph = document.add_paragraph()
        if bold_word and bold_word in text:
            before, _, after = text.partition(bold_word)
            paragraph.add_run(before)
            run = paragraph.add_run(bold_word)
            run.bold = True
            paragraph.add_run(after)
        elif tabs:
            paragraph.add_run(text)
            paragraph.add_run("\t")
            paragraph.add_run("tabbed")
        else:
            paragraph.add_run(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for col_index, value in enumerate(row):
                grid.cell(row_index, col_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    if tracked:
        data = _inject_tracked_changes(data, tracked)
    return data


def _inject_tracked_changes(data: bytes, tracked: dict) -> bytes:
    """Insert w:ins/w:del runs into the first paragraph of a real DOCX."""
    source = zipfile.ZipFile(io.BytesIO(data))
    document = source.read("word/document.xml").decode("utf-8")
    inserted = tracked.get("inserted", "")
    deleted = tracked.get("deleted", "")
    revision = (
        f'<w:ins w:id="1" w:author="Synthetic Author" w:date="2026-01-01T00:00:00Z">'
        f'<w:r><w:t>{inserted}</w:t></w:r></w:ins>'
        f'<w:del w:id="2" w:author="Synthetic Author" w:date="2026-01-01T00:00:00Z">'
        f'<w:r><w:delText>{deleted}</w:delText></w:r></w:del>'
    )
    document = document.replace("</w:p>", revision + "</w:p>", 1)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            payload = document.encode("utf-8") if item.filename == "word/document.xml" else source.read(item.filename)
            target.writestr(item, payload)
    source.close()
    return output.getvalue()


def build_pptx(*, slides=(("Slide one title", "Slide one body"),),
               notes=("Speaker note one",)) -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    for index, (title, body) in enumerate(slides):
        layout = presentation.slide_layouts[1]
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if index < len(notes):
            slide.notes_slide.notes_text_frame.text = notes[index]
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def build_xlsx(*, sheets=(("Sheet1", [["Header", "Value"], ["Row", 42]]),),
               formula=None) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets:
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    if formula:
        sheet_name, cell, expression = formula
        workbook[sheet_name][cell] = expression
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_native_pdf(text: str = "Native PDF text") -> bytes:
    """A minimal single-page PDF with a real text object (no OCR needed)."""
    content = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = io.BytesIO()
    output.write(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(output.tell())
        output.write(f"{index} 0 obj\n".encode())
        output.write(body)
        output.write(b"\nendobj\n")
    xref = output.tell()
    output.write(f"xref\n0 {len(objects) + 1}\n".encode())
    output.write(b"0000000000 65535 f \n")
    for offset in offsets:
        output.write(f"{offset:010d} 00000 n \n".encode())
    output.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
                 f"startxref\n{xref}\n%%EOF".encode())
    return output.getvalue()


def write(target: Path, data: bytes) -> Path:
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(data)
    return target


def build_text_png(text: str = "Synthetic text") -> bytes:
    """A real PNG with printed text for the image adapter's injected OCR."""
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (800, 200), "white")
    draw = ImageDraw.Draw(image)
    draw.text((40, 80), text, fill="black", font=ImageFont.load_default(size=40))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()
